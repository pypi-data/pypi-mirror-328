import re
from copy import deepcopy

from template_reports.templating import process_text, get_matching_tags

from .paragraphs import process_paragraph


def process_table_cell(cell, context, perm_user=None):
    """
    Process the text in a table cell.

    If the cell's text contains exactly one placeholder, process in "table" mode.
    If the result is a list, expand the table by adding new rows.
    Otherwise, process each paragraph in "normal" mode.
    """
    cell_text = cell.text.strip()

    matches = get_matching_tags(cell_text)

    # Process in "table" mode if exactly one tag is found.
    if len(matches) == 1:
        result = process_text(
            cell_text,
            context,
            perm_user=perm_user,
            mode="table",
        )
        if isinstance(result, list):
            expand_table_cell_with_list(cell, result)
        else:
            cell.text = result

    # Otherwise, process each paragraph in "normal" mode.
    else:
        for paragraph in cell.text_frame.paragraphs:
            process_paragraph(
                paragraph,
                context,
                perm_user,
                mode="normal",
            )


def clone_row_with_value(original_row, cell_index, new_text):
    """
    Helper function to create a new row from an existing row.

    - Deep-copies the original row.
    - Finds the cell at the given cell_index.
    - Updates that cell's text with new_text.
    - Returns the new row element.

    Uses _Cell to update the text properly.
    """

    from pptx.table import _Cell

    new_row = deepcopy(original_row)
    new_row_cells = list(new_row)
    if cell_index < len(new_row_cells):
        target_cell_element = new_row_cells[cell_index]
        new_cell = _Cell(target_cell_element, new_row)
        new_cell.text = str(new_text)
    return new_row


def expand_table_cell_with_list(cell, items):
    """
    Expand the table by adding a new row for each additional item in `items`.
    The original cell is updated with the first item; for each subsequent item, a new row is cloned
    and the corresponding cell (at the same column index) is updated with the item's text.

    If the cell's underlying XML does not have a parent (i.e. row or table not found), simply update the cell.
    """
    if not items:
        cell.text = ""
        return

    # Set first item in original cell.
    cell.text = str(items[0])

    # Attempt to get the row element (<a:tr>) from the cell's XML.
    row_element = cell._tc.getparent()
    if row_element is None:
        # Fallback: if row element is not available, do not expand.
        return
    table_element = row_element.getparent()
    if table_element is None:
        # Fallback: if table element is not available, do not expand.
        return

    # Determine the cell's index in the row.
    row_cells = list(row_element)
    cell_index = row_cells.index(cell._tc)

    # For each additional item, clone the row.
    for item in items[1:]:
        new_row_element = clone_row_with_value(row_element, cell_index, item)
        table_element.append(new_row_element)
